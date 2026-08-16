import re
from collections import Counter
from datetime import datetime
from io import BytesIO
from zipfile import BadZipFile

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import Candidate, CandidateStatus, RecruitmentStage
from app.schemas import ImportPreviewResult, ImportPreviewRow, ImportResult
from app.services.audit_service import AuditEvent, append_audit_event


HEADER_ERROR_CODE = "CANDIDATE_HEADER_INVALID"
REQUIRED_VALUE_ERROR_CODE = "CANDIDATE_REQUIRED_VALUE_MISSING"
DUPLICATE_APPLICATION_ID_ERROR_CODE = "DUPLICATE_APPLICATION_ID"
INVALID_FIELD_ERROR_CODE = "CANDIDATE_FIELD_INVALID"

HEADER_ALIASES = {
    "application_id": "application_id",
    "candidate_name": "full_name",
    "full_name": "full_name",
    "name": "full_name",
    "email": "email",
    "stage": "stage",
    "status": "status",
    "position": "position",
    "phone": "phone",
    "interview_time": "interview_time",
    "interviewer": "interviewer",
    "note": "note",
}
REQUIRED_HEADERS = {"application_id", "full_name", "email", "stage", "status"}
REQUIRED_FIELD_LABELS = {
    "application_id": "Application ID",
    "full_name": "Candidate Name",
    "email": "Email",
    "stage": "Stage",
    "status": "Status",
}
VALID_STAGES = {stage.value for stage in RecruitmentStage}
VALID_STATUSES = {
    CandidateStatus.PENDING.value,
    CandidateStatus.PASS_CV.value,
    CandidateStatus.REJECT_CV.value,
    CandidateStatus.PASS_INTERVIEW.value,
    CandidateStatus.REJECT_INTERVIEW.value,
}
EMAIL_PATTERN = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
FIELD_MAX_LENGTHS = {
    "application_id": 64,
    "full_name": 255,
    "email": 255,
    "phone": 50,
    "position": 255,
    "interviewer": 255,
}


async def preview_candidates_from_excel(
    db: Session,
    file: UploadFile,
) -> ImportPreviewResult:
    content = await read_excel_content(file)
    preview_rows = build_preview_rows(db, content)
    valid_rows = [row for row in preview_rows if row.is_valid]

    return ImportPreviewResult(
        total_rows=len(preview_rows),
        valid_rows=len(valid_rows),
        invalid_rows=len(preview_rows) - len(valid_rows),
        rows=preview_rows,
    )


async def import_candidates_from_excel(
    db: Session,
    file: UploadFile,
) -> ImportResult:
    content = await read_excel_content(file)
    preview_rows = build_preview_rows(db, content)
    imported_count = 0
    errors: list[str] = []

    for row in preview_rows:
        if not row.is_valid:
            errors.append(f"Row {row.row_number}: {row.reason}")
            continue

        candidate = create_candidate_from_payload(row.candidate)
        db.add(candidate)
        db.flush()
        append_audit_event(
            db,
            event=AuditEvent.CANDIDATE_IMPORTED,
            entity_type="CANDIDATE",
            entity_id=candidate.id,
            application_id=candidate.application_id,
            actor="demo_hr",
            payload={
                "row_number": row.row_number,
                "stage": candidate.stage,
                "status": candidate.status,
            },
        )
        imported_count += 1

    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        if "application_id" in str(error.orig).lower():
            raise HTTPException(
                status_code=409,
                detail=(
                    f"{DUPLICATE_APPLICATION_ID_ERROR_CODE}: An Application ID was "
                    "created by another request during this import. Preview the file again."
                ),
            ) from error
        raise

    return ImportResult(
        imported=imported_count,
        skipped=len(preview_rows) - imported_count,
        errors=errors,
    )


async def read_excel_content(file: UploadFile) -> bytes:
    if file.filename and not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload a .xlsx Excel file.")

    return await file.read()


def build_preview_rows(
    db: Session,
    content: bytes,
) -> list[ImportPreviewRow]:
    workbook = open_workbook(content)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)

    if not header_row or all(cell.value is None for cell in header_row):
        raise HTTPException(status_code=400, detail="Excel file is empty.")

    headers = [str(cell.value).strip() if cell.value else "" for cell in header_row]
    normalized_headers = validate_and_normalize_headers(headers)
    source_rows: list[tuple[int, dict[str, object]]] = []

    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not is_empty_row(row):
            source_rows.append(
                (row_index, build_row_payload(normalized_headers, row))
            )

    application_ids = [
        application_id
        for _, payload in source_rows
        if (application_id := to_optional_string(payload.get("application_id")))
    ]
    application_id_counts = Counter(application_ids)
    duplicate_file_ids = {
        application_id
        for application_id, count in application_id_counts.items()
        if count > 1
    }
    existing_application_ids = {
        application_id
        for (application_id,) in (
            db.query(Candidate.application_id)
            .filter(Candidate.application_id.in_(set(application_ids)))
            .all()
        )
    }
    preview_rows: list[ImportPreviewRow] = []

    for row_index, raw_payload in source_rows:
        candidate_payload, errors = validate_candidate_payload(raw_payload)
        application_id = candidate_payload.get("application_id")

        if application_id in existing_application_ids:
            errors.append(
                f"{DUPLICATE_APPLICATION_ID_ERROR_CODE}: Application ID already exists: "
                f"{application_id}"
            )
        elif application_id in duplicate_file_ids:
            errors.append(
                f"{DUPLICATE_APPLICATION_ID_ERROR_CODE}: Application ID is repeated in this "
                f"file: {application_id}"
            )

        preview_rows.append(
            ImportPreviewRow(
                row_number=row_index,
                is_valid=len(errors) == 0,
                reason="; ".join(errors) if errors else None,
                candidate=candidate_payload,
            )
        )

    return preview_rows


def open_workbook(content: bytes):
    try:
        return load_workbook(BytesIO(content))
    except (BadZipFile, InvalidFileException) as error:
        raise HTTPException(status_code=400, detail="Invalid Excel file format.") from error


def validate_candidate_payload(
    payload: dict[str, object],
) -> tuple[dict[str, object], list[str]]:
    errors: list[str] = []
    application_id = to_optional_string(payload.get("application_id"))
    full_name = to_optional_string(payload.get("full_name"))
    email = normalize_email(payload.get("email"), errors)
    stage = normalize_stage(payload.get("stage"))
    status = normalize_status(payload.get("status"))
    interview_time = None

    required_values = {
        "application_id": application_id,
        "full_name": full_name,
        "email": to_optional_string(payload.get("email")),
        "stage": to_optional_string(payload.get("stage")),
        "status": to_optional_string(payload.get("status")),
    }
    missing_fields = [
        REQUIRED_FIELD_LABELS[field_name]
        for field_name, value in required_values.items()
        if value is None
    ]
    if missing_fields:
        errors.append(
            f"{REQUIRED_VALUE_ERROR_CODE}: Missing required value(s): "
            f"{', '.join(missing_fields)}"
        )

    for field_name, max_length in FIELD_MAX_LENGTHS.items():
        field_value = to_optional_string(payload.get(field_name))
        if field_value and len(field_value) > max_length:
            errors.append(
                f"{INVALID_FIELD_ERROR_CODE}: {field_name} exceeds {max_length} characters"
            )

    if to_optional_string(payload.get("stage")) and stage is None:
        errors.append(
            f"{INVALID_FIELD_ERROR_CODE}: Unsupported Stage: {payload.get('stage')}"
        )
    if to_optional_string(payload.get("status")) and status is None:
        errors.append(
            f"{INVALID_FIELD_ERROR_CODE}: Unsupported Status: {payload.get('status')}"
        )

    try:
        interview_time = to_optional_datetime(payload.get("interview_time"))
    except ValueError:
        errors.append(
            f"{INVALID_FIELD_ERROR_CODE}: Invalid Interview Time: "
            f"{payload.get('interview_time')}"
        )

    candidate_payload: dict[str, object] = {
        "application_id": application_id or "",
        "full_name": full_name or "",
        "email": email or "",
        "phone": to_optional_string(payload.get("phone")),
        "position": to_optional_string(payload.get("position")),
        "stage": stage or "",
        "status": status or "",
        "interview_time": interview_time,
        "interviewer": to_optional_string(payload.get("interviewer")),
        "note": to_optional_string(payload.get("note")),
    }

    return candidate_payload, errors


def build_row_payload(
    headers: list[str],
    row: tuple[object, ...],
) -> dict[str, object]:
    return {
        headers[index]: row[index]
        for index in range(min(len(headers), len(row)))
        if headers[index]
    }


def create_candidate_from_payload(
    payload: dict[str, object],
) -> Candidate:
    return Candidate(
        application_id=str(payload.get("application_id")).strip(),
        full_name=str(payload.get("full_name")).strip(),
        email=str(payload.get("email")).strip().lower(),
        phone=to_optional_string(payload.get("phone")),
        position=to_optional_string(payload.get("position")),
        stage=normalize_stage(payload.get("stage")) or "",
        status=normalize_status(payload.get("status")) or "",
        interview_time=to_optional_datetime(payload.get("interview_time")),
        interviewer=to_optional_string(payload.get("interviewer")),
        note=to_optional_string(payload.get("note")),
    )


def normalize_header(
    header: str,
) -> str:
    normalized = header.strip().lower().replace("-", "_").replace(" ", "_")
    return HEADER_ALIASES.get(normalized, normalized)


def validate_and_normalize_headers(headers: list[str]) -> list[str]:
    normalized_headers = [normalize_header(header) for header in headers]
    present_headers = {header for header in normalized_headers if header}
    missing_headers = sorted(REQUIRED_HEADERS - present_headers)
    duplicate_headers = sorted(
        header
        for header in present_headers
        if normalized_headers.count(header) > 1
    )

    if missing_headers or duplicate_headers:
        details: list[str] = []
        if missing_headers:
            details.append(
                "missing "
                + ", ".join(REQUIRED_FIELD_LABELS[name] for name in missing_headers)
            )
        if duplicate_headers:
            details.append("duplicate " + ", ".join(duplicate_headers))
        raise HTTPException(
            status_code=400,
            detail=f"{HEADER_ERROR_CODE}: Invalid Excel header contract ({'; '.join(details)}).",
        )

    return normalized_headers


def normalize_stage(value: object) -> str | None:
    raw_stage = to_optional_string(value)
    if not raw_stage:
        return None

    normalized_stage = normalize_enum_value(raw_stage)
    return normalized_stage if normalized_stage in VALID_STAGES else None


def normalize_status(value: object) -> str | None:
    raw_status = to_optional_string(value)

    if not raw_status:
        return None

    normalized_status = normalize_enum_value(raw_status)

    return normalized_status if normalized_status in VALID_STATUSES else None


def normalize_enum_value(value: str) -> str:
    return value.strip().upper().replace("-", "_").replace(" ", "_")


def normalize_email(value: object, errors: list[str]) -> str | None:
    raw_email = to_optional_string(value)
    if not raw_email:
        return None

    if EMAIL_PATTERN.fullmatch(raw_email) is None:
        errors.append(f"{INVALID_FIELD_ERROR_CODE}: Invalid Email: {raw_email}")
        return None

    return raw_email.lower()


def to_optional_string(
    value: object,
) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def to_optional_datetime(
    value: object,
) -> datetime | None:
    if isinstance(value, datetime):
        return value

    if isinstance(value, str) and value.strip():
        return datetime.fromisoformat(value.strip())

    return None


def is_empty_row(row: tuple[object, ...]) -> bool:
    return all(value is None or str(value).strip() == "" for value in row)
